from __future__ import annotations

from collections.abc import Iterator
from datetime import date, datetime
import csv
import hashlib
import json
from pathlib import Path, PurePosixPath
import re
from xml.etree import ElementTree
import zipfile

from MetadataArea.importing.domain import MetadataImportInterpretation, MetadataImportProbe


SUPPORTED_SUFFIXES = (
    ".csv", ".tsv", ".tab", ".txt", ".json", ".jsonl", ".ndjson",
    ".xml", ".tei", ".xlsx", ".ods", ".parquet", ".arrow", ".ipc", ".feather", ".zip",
)
IMPORT_FILE_FILTER = (
    "Metadata files (*.csv *.tsv *.tab *.txt *.json *.jsonl *.ndjson *.xml *.tei "
    "*.xlsx *.ods *.parquet *.arrow *.ipc *.feather *.zip);;"
    "Delimited text (*.csv *.tsv *.tab *.txt);;JSON (*.json *.jsonl *.ndjson);;"
    "XML / TEI (*.xml *.tei *.zip);;Spreadsheets (*.xlsx *.ods);;"
    "Columnar tables (*.parquet *.arrow *.ipc *.feather);;All files (*)"
)

_XML_NAMESPACE = "{http://www.w3.org/XML/1998/namespace}"
_MAX_ARCHIVE_MEMBERS = 10_000
_MAX_ARCHIVE_EXPANDED_BYTES = 2 * 1024 * 1024 * 1024
_MAX_ARCHIVE_RATIO = 500


def source_fingerprint(path: Path) -> str:
    path = Path(path)
    digest = hashlib.sha256()
    if path.is_dir():
        files = sorted(item for item in path.rglob("*") if item.is_file())
        for item in files:
            stat = item.stat()
            digest.update(item.relative_to(path).as_posix().encode("utf-8"))
            digest.update(str(stat.st_size).encode("ascii"))
            digest.update(str(stat.st_mtime_ns).encode("ascii"))
        return digest.hexdigest()
    stat = path.stat()
    digest.update(str(path.resolve()).encode("utf-8"))
    digest.update(str(stat.st_size).encode("ascii"))
    digest.update(str(stat.st_mtime_ns).encode("ascii"))
    return digest.hexdigest()


def probe_source(path: Path) -> MetadataImportProbe:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    if path.is_dir():
        xml_files = sorted((*path.rglob("*.xml"), *path.rglob("*.tei")))
        if not xml_files:
            raise ValueError("The selected folder does not contain XML or TEI files.")
        return MetadataImportProbe(
            source_path=path,
            format_key="tei_collection",
            format_label="TEI / XML folder",
            size_bytes=sum(item.stat().st_size for item in xml_files),
            source_fingerprint=source_fingerprint(path),
            row_count_hint=len(xml_files),
            default_record_path="TEI",
            is_tei=True,
        )
    if not path.is_file() or path.stat().st_size <= 0:
        raise ValueError("Choose a nonempty readable metadata source.")

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported metadata format: {suffix or 'file without an extension'}")
    size = path.stat().st_size
    fingerprint = source_fingerprint(path)
    columns: tuple[str, ...] = ()
    tables: tuple[str, ...] = ()
    default_table = ""
    row_hint = None
    record_path = ""
    is_tei = False
    warnings: list[str] = []
    default_encoding = "utf-8-sig"
    encoding_ambiguous = False
    default_delimiter = ""
    delimiter_ambiguous = False

    if suffix in {".csv", ".tsv", ".tab", ".txt"}:
        default_encoding, encoding_ambiguous = _detect_text_encoding(path)
        if suffix in {".tsv", ".tab"}:
            delimiter = "\t"
        else:
            delimiter, delimiter_ambiguous = _sniff_delimiter_details(path, default_encoding)
        default_delimiter = delimiter
        with path.open("r", encoding=default_encoding, newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter)
            columns = tuple(str(value).strip() for value in next(reader, ()))
        if not columns:
            raise ValueError("Delimited metadata is missing a header row.")
        if encoding_ambiguous:
            warnings.append("Text encoding was inferred as Windows-1252; verify it before importing.")
        format_key, label = "delimited", "Delimited text"
    elif suffix in {".json", ".jsonl", ".ndjson"}:
        format_key = "jsonl" if suffix != ".json" else "json"
        label = "JSON Lines" if format_key == "jsonl" else "JSON"
        first, record_path = _probe_json(path, lines=format_key == "jsonl")
        if first is not None:
            columns = tuple(_flatten_record(first, "; ").keys())
    elif suffix in {".xml", ".tei"}:
        root = _safe_xml_root(path.read_bytes(), path.name)
        root_name = _local_name(root.tag)
        is_tei = root_name in {"TEI", "teiCorpus"} or suffix == ".tei"
        format_key = "tei" if is_tei else "xml"
        label = "TEI P5 XML" if is_tei else "XML"
        record_path = "TEI" if is_tei else _default_xml_record_name(root)
    elif suffix == ".zip":
        members = _validated_xml_members(path)
        first_bytes = _read_zip_member(path, members[0])
        first_root = _safe_xml_root(first_bytes, members[0].filename)
        is_tei = _local_name(first_root.tag) in {"TEI", "teiCorpus"}
        format_key = "tei_zip" if is_tei else "xml_zip"
        label = "TEI ZIP collection" if is_tei else "XML ZIP collection"
        row_hint = len(members)
        record_path = "TEI" if is_tei else _default_xml_record_name(first_root)
    elif suffix == ".xlsx":
        format_key, label = "xlsx", "Excel workbook"
        tables = _spreadsheet_names(path, ods=False)
        default_table = tables[0] if tables else ""
    elif suffix == ".ods":
        format_key, label = "ods", "OpenDocument spreadsheet"
        tables = _spreadsheet_names(path, ods=True)
        default_table = tables[0] if tables else ""
    elif suffix == ".parquet":
        format_key, label = "parquet", "Apache Parquet"
        columns, row_hint = _probe_arrow_table(path, parquet=True)
    else:
        format_key, label = "arrow", "Apache Arrow IPC"
        columns, row_hint = _probe_arrow_table(path, parquet=False)

    return MetadataImportProbe(
        source_path=path,
        format_key=format_key,
        format_label=label,
        size_bytes=size,
        source_fingerprint=fingerprint,
        row_count_hint=row_hint,
        column_names=columns,
        table_names=tables,
        default_table=default_table,
        default_record_path=record_path,
        is_tei=is_tei,
        default_encoding=default_encoding,
        encoding_ambiguous=encoding_ambiguous,
        default_delimiter=default_delimiter,
        delimiter_ambiguous=delimiter_ambiguous,
        warnings=tuple(warnings),
    )


def iter_records(
    probe: MetadataImportProbe,
    interpretation: MetadataImportInterpretation,
) -> Iterator[dict[str, str]]:
    key = probe.format_key
    path = probe.source_path
    separator = interpretation.repeated_value_separator or "; "
    if key == "delimited":
        yield from _iter_delimited(path, interpretation)
    elif key in {"json", "jsonl"}:
        yield from _iter_json(path, interpretation, lines=key == "jsonl")
    elif key in {"xml", "tei"}:
        yield from _records_from_xml_bytes(
            path.read_bytes(), path.name, interpretation, force_tei=key == "tei"
        )
    elif key == "tei_collection":
        for source in sorted((*path.rglob("*.xml"), *path.rglob("*.tei"))):
            yield from _records_from_xml_bytes(
                source.read_bytes(), source.name, interpretation, force_tei=True
            )
    elif key in {"tei_zip", "xml_zip"}:
        for member in _validated_xml_members(path):
            yield from _records_from_xml_bytes(
                _read_zip_member(path, member),
                member.filename,
                interpretation,
                force_tei=key == "tei_zip",
            )
    elif key == "xlsx":
        yield from _iter_xlsx(path, interpretation)
    elif key == "ods":
        yield from _iter_ods(path, interpretation)
    elif key in {"parquet", "arrow"}:
        yield from _iter_arrow(path, parquet=key == "parquet", separator=separator)
    else:
        raise ValueError(f"Unsupported metadata reader: {key}")


def _iter_delimited(path: Path, options: MetadataImportInterpretation) -> Iterator[dict[str, str]]:
    delimiter = options.delimiter or ("\t" if path.suffix.lower() in {".tsv", ".tab"} else _sniff_delimiter(path, options.encoding))
    with path.open("r", encoding=options.encoding or "utf-8-sig", newline="") as handle:
        for _ in range(max(0, options.header_row - 1)):
            next(handle, None)
        reader = csv.DictReader(handle, delimiter=delimiter)
        if not reader.fieldnames:
            raise ValueError("Delimited metadata is missing a header row.")
        names = _normalized_names(reader.fieldnames)
        for raw in reader:
            values = list(raw.values())
            if len(values) > len(names):
                raise ValueError("A delimited row contains more cells than the header.")
            yield {
                name: _scalar_text(values[index] if index < len(values) else "")
                for index, name in enumerate(names)
            }


def _iter_json(path: Path, options: MetadataImportInterpretation, *, lines: bool) -> Iterator[dict[str, str]]:
    separator = options.repeated_value_separator or "; "
    if lines:
        with path.open("r", encoding=options.encoding or "utf-8-sig") as handle:
            for line_number, raw_line in enumerate(handle, 1):
                if not raw_line.strip():
                    continue
                try:
                    value = json.loads(raw_line)
                except json.JSONDecodeError as exc:
                    raise ValueError(f"Invalid JSONL record at line {line_number}: {exc}") from exc
                yield _flatten_record(value, separator)
        return

    with path.open("r", encoding=options.encoding or "utf-8-sig") as handle:
        prefix = handle.read(4096).lstrip()
    if prefix.startswith("[") or options.record_path:
        try:
            import ijson
        except ImportError:
            pass
        else:
            item_path = f"{options.record_path}.item" if options.record_path else "item"
            with path.open("rb") as handle:
                found = False
                for record in ijson.items(handle, item_path):
                    found = True
                    yield _flatten_record(record, separator)
            if found:
                return
    payload = json.loads(path.read_text(encoding=options.encoding or "utf-8-sig"))
    records = _json_record_collection(payload, options.record_path)
    for record in records:
        yield _flatten_record(record, separator)


def _records_from_xml_bytes(
    raw: bytes,
    source_name: str,
    options: MetadataImportInterpretation,
    *,
    force_tei: bool,
) -> Iterator[dict[str, str]]:
    root = _safe_xml_root(raw, source_name)
    is_tei = force_tei or _local_name(root.tag) in {"TEI", "teiCorpus"}
    profile = "tei_p5" if is_tei else "generic"
    yield from XML_ADAPTER_REGISTRY[profile](root, source_name, options)


def _iter_tei_xml_records(
    root: ElementTree.Element,
    source_name: str,
    options: MetadataImportInterpretation,
) -> Iterator[dict[str, str]]:
    elements = (
        [root]
        if _local_name(root.tag) == "TEI"
        else [item for item in root.iter() if _local_name(item.tag) == "TEI"]
    )
    for element in elements:
        yield _tei_record(
            element,
            source_name,
            options.repeated_value_separator or "; ",
        )


def _iter_generic_xml_records(
    root: ElementTree.Element,
    _source_name: str,
    options: MetadataImportInterpretation,
) -> Iterator[dict[str, str]]:

    record_name = (options.record_path or _default_xml_record_name(root)).strip()
    records = [item for item in root.iter() if _local_name(item.tag) == record_name]
    if not records and _local_name(root.tag) == record_name:
        records = [root]
    if not records:
        raise ValueError(f"XML does not contain record elements named {record_name!r}.")
    for element in records:
        yield _flatten_xml_record(element, options.repeated_value_separator or "; ")


# Explicit profile registry keeps the staged record contract reusable for
# future Dublin Core, MODS, EAD, and METS adapters.
XML_ADAPTER_REGISTRY = {
    "generic": _iter_generic_xml_records,
    "tei_p5": _iter_tei_xml_records,
}


def _tei_record(element: ElementTree.Element, source_name: str, separator: str) -> dict[str, str]:
    header = next((item for item in element if _local_name(item.tag) == "teiHeader"), None)
    if header is None:
        header = element

    def texts(*paths: tuple[str, ...]) -> list[str]:
        output: list[str] = []
        for item in header.iter():
            name = _local_name(item.tag)
            if any(name == path[-1] and _has_ancestor_path(header, item, path[:-1]) for path in paths):
                value = _element_text(item)
                if value and value not in output:
                    output.append(value)
        return output

    xml_id = str(element.attrib.get(f"{_XML_NAMESPACE}id", "")).strip()
    idnos = texts(("publicationStmt", "idno"), ("sourceDesc", "idno"))
    languages: list[str] = []
    publication_dates: list[str] = []
    licences: list[str] = []
    classifications: list[str] = []
    revisions: list[str] = []
    for item in header.iter():
        item_name = _local_name(item.tag)
        if item_name == "language":
            value = str(item.attrib.get("ident") or _element_text(item)).strip()
            if value and value not in languages:
                languages.append(value)
        elif item_name == "date" and _has_ancestor_path(header, item, ("publicationStmt",)):
            value = str(item.attrib.get("when") or _element_text(item)).strip()
            if value and value not in publication_dates:
                publication_dates.append(value)
        elif item_name == "licence" and _has_ancestor_path(header, item, ("availability",)):
            value = str(item.attrib.get("target") or _element_text(item)).strip()
            if value and value not in licences:
                licences.append(value)
        elif item_name in {"classCode", "catRef"} and _has_ancestor_path(header, item, ("textClass",)):
            value = str(item.attrib.get("target") or item.attrib.get("scheme") or _element_text(item)).strip()
            if value and value not in classifications:
                classifications.append(value)
        elif item_name == "change" and _has_ancestor_path(header, item, ("revisionDesc",)):
            value = str(item.attrib.get("when") or _element_text(item)).strip()
            if value:
                revisions.append(value)
    record = {
        "source_filename": PurePosixPath(source_name).name,
        "identifier": xml_id or (idnos[0] if idnos else PurePosixPath(source_name).stem),
        "title": separator.join(texts(("titleStmt", "title"))),
        "authors": separator.join(texts(("titleStmt", "author"))),
        "publication_date": separator.join(publication_dates),
        "publisher": separator.join(texts(("publicationStmt", "publisher"))),
        "publication_place": separator.join(texts(("publicationStmt", "pubPlace"))),
        "licence": separator.join(licences),
        "source_description": separator.join(texts(("sourceDesc", "bibl"), ("sourceDesc", "p"))),
        "languages": separator.join(languages),
        "keywords": separator.join(texts(("keywords", "term"))),
        "classifications": separator.join(classifications),
        "latest_revision": max(revisions) if revisions else "",
    }
    return record


def _has_ancestor_path(root: ElementTree.Element, target: ElementTree.Element, path: tuple[str, ...]) -> bool:
    if not path:
        return True
    # ElementTree does not expose parents. Walking short TEI header paths keeps the
    # adapter deterministic without enabling arbitrary XPath execution.
    def visit(node: ElementTree.Element, ancestors: tuple[str, ...]) -> bool:
        current = (*ancestors, _local_name(node.tag))
        if node is target:
            return len(current) >= len(path) + 1 and current[-len(path)-1:-1] == path
        return any(visit(child, current) for child in node)
    return visit(root, ())


def _flatten_xml_record(element: ElementTree.Element, separator: str) -> dict[str, str]:
    values: dict[str, list[str]] = {}
    for key, value in element.attrib.items():
        values.setdefault(f"@{_local_name(key)}", []).append(str(value).strip())

    def walk(node: ElementTree.Element, prefix: str) -> None:
        children = list(node)
        name = _local_name(node.tag)
        current = f"{prefix}.{name}" if prefix else name
        for key, value in node.attrib.items():
            values.setdefault(f"{current}.@{_local_name(key)}", []).append(str(value).strip())
        if not children:
            values.setdefault(current, []).append(_element_text(node))
        else:
            for child in children:
                walk(child, current)

    for child in element:
        walk(child, "")
    if not values and _element_text(element):
        values["value"] = [_element_text(element)]
    return {name: separator.join(item for item in items if item) for name, items in values.items()}


def _iter_xlsx(path: Path, options: MetadataImportInterpretation) -> Iterator[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Excel import requires the openpyxl package.") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[options.table_name or workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        for _ in range(max(0, options.header_row - 1)):
            next(rows, None)
        names = _normalized_names(next(rows, ()))
        if not names:
            raise ValueError("The selected spreadsheet sheet has no header row.")
        for values in rows:
            yield {name: _scalar_text(values[index] if index < len(values) else "") for index, name in enumerate(names)}
    finally:
        workbook.close()


def _iter_ods(path: Path, options: MetadataImportInterpretation) -> Iterator[dict[str, str]]:
    try:
        from odf.opendocument import load
        from odf.table import Table, TableCell, TableRow
        from odf.text import P
    except ImportError as exc:
        raise ValueError("ODS import requires the odfpy package.") from exc
    document = load(str(path))
    tables = document.spreadsheet.getElementsByType(Table)
    selected = options.table_name or (str(tables[0].getAttribute("name")) if tables else "")
    table = next((item for item in tables if str(item.getAttribute("name")) == selected), None)
    if table is None:
        raise ValueError("Choose a spreadsheet sheet to import.")
    matrix: list[list[str]] = []
    for row in table.getElementsByType(TableRow):
        values: list[str] = []
        for cell in row.getElementsByType(TableCell):
            value_type = str(cell.getAttribute("valuetype") or "")
            if value_type == "date":
                text = str(cell.getAttribute("datevalue") or "")
            elif value_type == "boolean":
                text = "true" if str(cell.getAttribute("booleanvalue")).casefold() == "true" else "false"
            elif value_type in {"float", "currency", "percentage"}:
                text = str(cell.getAttribute("value") or "")
            else:
                text = "".join(
                    str(node)
                    for paragraph in cell.getElementsByType(P)
                    for node in paragraph.childNodes
                )
            repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
            values.extend([text] * min(repeat, 10_000))
        row_repeat = int(row.getAttribute("numberrowsrepeated") or 1)
        matrix.extend([values] * min(row_repeat, 10_000))
    start = max(0, options.header_row - 1)
    if start >= len(matrix):
        raise ValueError("The selected spreadsheet sheet has no header row.")
    names = _normalized_names(matrix[start])
    for values in matrix[start + 1:]:
        yield {name: _scalar_text(values[index] if index < len(values) else "") for index, name in enumerate(names)}


def _iter_arrow(path: Path, *, parquet: bool, separator: str) -> Iterator[dict[str, str]]:
    try:
        import pyarrow.ipc as ipc
        import pyarrow.parquet as pq
    except ImportError as exc:
        raise ValueError("Parquet and Arrow import require the pyarrow package.") from exc
    if parquet:
        source = pq.ParquetFile(path).iter_batches(batch_size=4096)
    else:
        handle = path.open("rb")
        try:
            try:
                reader = ipc.open_file(handle)
                source = (reader.get_batch(index) for index in range(reader.num_record_batches))
            except Exception:
                handle.seek(0)
                source = ipc.open_stream(handle)
            for batch in source:
                for record in batch.to_pylist():
                    yield _flatten_record(record, separator)
        finally:
            handle.close()
        return
    for batch in source:
        for record in batch.to_pylist():
            yield _flatten_record(record, separator)


def _probe_json(path: Path, *, lines: bool) -> tuple[dict[str, object] | None, str]:
    if lines:
        with path.open("r", encoding="utf-8-sig") as handle:
            for raw in handle:
                if raw.strip():
                    value = json.loads(raw)
                    if not isinstance(value, dict):
                        raise ValueError("Each JSONL line must contain an object.")
                    return value, ""
        raise ValueError("The JSONL source contains no records.")
    try:
        import ijson
    except ImportError:
        pass
    else:
        with path.open("rb") as handle:
            first_character = handle.read(4096).lstrip()[:1]
        if first_character == b"[":
            with path.open("rb") as handle:
                first = next(ijson.items(handle, "item"), None)
            if first is not None and not isinstance(first, dict):
                raise ValueError("The selected JSON record collection must contain objects.")
            return first, ""
        if first_character == b"{":
            record_path = ""
            with path.open("rb") as handle:
                for prefix, event, _value in ijson.parse(handle):
                    if event == "start_map" and prefix.endswith(".item"):
                        record_path = prefix[:-5]
                        break
            if record_path:
                with path.open("rb") as handle:
                    first = next(ijson.items(handle, f"{record_path}.item"), None)
                if first is not None and not isinstance(first, dict):
                    raise ValueError("The selected JSON record collection must contain objects.")
                return first, record_path
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(payload, list):
        if payload and not isinstance(payload[0], dict):
            raise ValueError("The selected JSON record collection must contain objects.")
        return (payload[0] if payload else None), ""
    if isinstance(payload, dict):
        list_keys = [key for key, value in payload.items() if isinstance(value, list) and value and isinstance(value[0], dict)]
        if list_keys:
            record_path = str(list_keys[0])
            return payload[list_keys[0]][0], record_path
        return payload, ""
    raise ValueError("JSON metadata must contain an object or an array of objects.")


def _json_record_collection(payload: object, record_path: str) -> list[dict[str, object]]:
    value = payload
    if record_path:
        for part in record_path.split("."):
            if not isinstance(value, dict) or part not in value:
                raise ValueError(f"JSON record path {record_path!r} was not found.")
            value = value[part]
    if isinstance(value, list):
        if not all(isinstance(item, dict) for item in value):
            raise ValueError("The selected JSON record collection must contain objects.")
        return list(value)
    if isinstance(value, dict):
        candidates = [item for item in value.values() if isinstance(item, list) and all(isinstance(record, dict) for record in item)]
        if len(candidates) == 1:
            return list(candidates[0])
        return [value]
    raise ValueError("JSON metadata must contain an object or an array of objects.")


def _flatten_record(record: object, separator: str, prefix: str = "") -> dict[str, str]:
    if not isinstance(record, dict):
        raise ValueError("Metadata records must be objects.")
    output: dict[str, str] = {}
    for raw_key, value in record.items():
        key = str(raw_key).strip()
        name = f"{prefix}.{key}" if prefix else key
        if isinstance(value, dict):
            output.update(_flatten_record(value, separator, name))
        elif isinstance(value, list):
            if any(isinstance(item, (dict, list)) for item in value):
                raise ValueError(f"Nested object array {name!r} must be selected as the record collection.")
            if any(isinstance(item, (bytes, bytearray, memoryview)) for item in value):
                raise ValueError(f"Unsupported binary field {name!r}.")
            output[name] = separator.join(_scalar_text(item) for item in value if item is not None)
        elif isinstance(value, (bytes, bytearray, memoryview)):
            raise ValueError(f"Unsupported binary field {name!r}.")
        else:
            output[name] = _scalar_text(value)
    return output


def _scalar_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _normalized_names(values) -> tuple[str, ...]:
    names: list[str] = []
    used: set[str] = set()
    for index, value in enumerate(values, 1):
        base = str(value or "").strip()
        if not base:
            base = f"column_{index}"
        if base.lower() == "id":
            base = "id_imported"
        candidate = base
        suffix = 2
        while candidate.casefold() in used:
            candidate = f"{base}_{suffix}"
            suffix += 1
        used.add(candidate.casefold())
        names.append(candidate)
    return tuple(names)


def _sniff_delimiter(path: Path, encoding: str = "utf-8-sig") -> str:
    return _sniff_delimiter_details(path, encoding)[0]


def _sniff_delimiter_details(path: Path, encoding: str = "utf-8-sig") -> tuple[str, bool]:
    with path.open("r", encoding=encoding, newline="") as handle:
        sample = handle.read(8192)
    if not sample.strip():
        return ",", True
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter, False
    except csv.Error:
        return ",", True


def _detect_text_encoding(path: Path) -> tuple[str, bool]:
    prefix = path.read_bytes()[:131072]
    if prefix.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig", False
    if prefix.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16", False
    try:
        prefix.decode("utf-8")
    except UnicodeDecodeError:
        # Windows-1252 is a useful, deterministic fallback for Western legacy
        # tables. The UI exposes the choice because this branch is ambiguous.
        return "cp1252", True
    return "utf-8-sig", False


def _safe_xml_root(raw: bytes, source_name: str) -> ElementTree.Element:
    prefix = raw[:131072].upper()
    if b"<!DOCTYPE" in prefix or b"<!ENTITY" in prefix:
        raise ValueError(f"XML source {source_name!r} contains a prohibited DTD or entity declaration.")
    try:
        from lxml import etree
    except ImportError:
        try:
            return ElementTree.fromstring(raw)
        except ElementTree.ParseError as exc:
            raise ValueError(f"Invalid XML metadata in {source_name!r}: {exc}") from exc
    try:
        parser = etree.XMLParser(
            resolve_entities=False,
            no_network=True,
            load_dtd=False,
            huge_tree=False,
            recover=False,
        )
        return etree.fromstring(raw, parser=parser)
    except etree.XMLSyntaxError as exc:
        raise ValueError(f"Invalid XML metadata in {source_name!r}: {exc}") from exc


def _default_xml_record_name(root: ElementTree.Element) -> str:
    children = list(root)
    if not children:
        return _local_name(root.tag)
    if len(children) == 1 and list(children[0]):
        return _local_name(list(children[0])[0].tag)
    return _local_name(children[0].tag)


def _local_name(tag: object) -> str:
    return str(tag).rsplit("}", 1)[-1]


def _element_text(element: ElementTree.Element) -> str:
    return re.sub(r"\s+", " ", " ".join(element.itertext())).strip()


def _validated_xml_members(path: Path) -> list[zipfile.ZipInfo]:
    try:
        archive = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise ValueError("The selected ZIP is not a readable archive.") from exc
    with archive:
        all_members = [item for item in archive.infolist() if not item.is_dir()]
        if len(all_members) > _MAX_ARCHIVE_MEMBERS:
            raise ValueError("The selected ZIP contains too many files.")
        expanded = 0
        for item in all_members:
            parts = PurePosixPath(item.filename).parts
            if PurePosixPath(item.filename).is_absolute() or ".." in parts:
                raise ValueError("The selected ZIP contains an unsafe path.")
            if item.flag_bits & 0x1:
                raise ValueError("Encrypted ZIP entries are not supported.")
            expanded += item.file_size
            if expanded > _MAX_ARCHIVE_EXPANDED_BYTES:
                raise ValueError("The selected ZIP expands beyond the import safety limit.")
            if item.compress_size and item.file_size / item.compress_size > _MAX_ARCHIVE_RATIO:
                raise ValueError("The selected ZIP contains a suspiciously compressed entry.")
        members = [
            item for item in all_members
            if PurePosixPath(item.filename).suffix.lower() in {".xml", ".tei"}
        ]
        if not members:
            raise ValueError("The selected ZIP does not contain XML or TEI files.")
        return members


def _read_zip_member(path: Path, member: zipfile.ZipInfo) -> bytes:
    with zipfile.ZipFile(path) as archive:
        with archive.open(member) as handle:
            return handle.read(member.file_size + 1)


def _spreadsheet_names(path: Path, *, ods: bool) -> tuple[str, ...]:
    if ods:
        try:
            from odf.opendocument import load
            from odf.table import Table
        except ImportError:
            return ()
        document = load(str(path))
        return tuple(str(item.getAttribute("name")) for item in document.spreadsheet.getElementsByType(Table))
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return tuple(workbook.sheetnames)
    finally:
        workbook.close()


def _probe_arrow_table(path: Path, *, parquet: bool) -> tuple[tuple[str, ...], int | None]:
    try:
        import pyarrow.ipc as ipc
        import pyarrow.parquet as pq
    except ImportError as exc:
        raise ValueError("Parquet and Arrow import require the pyarrow package.") from exc
    if parquet:
        source = pq.ParquetFile(path)
        return tuple(source.schema_arrow.names), int(source.metadata.num_rows)
    with path.open("rb") as handle:
        try:
            reader = ipc.open_file(handle)
            return tuple(reader.schema.names), sum(reader.get_batch(index).num_rows for index in range(reader.num_record_batches))
        except Exception:
            handle.seek(0)
            reader = ipc.open_stream(handle)
            return tuple(reader.schema.names), None
